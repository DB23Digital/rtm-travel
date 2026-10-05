"""Build the editable .docx twin of the printable policy template.

The HTML at public/downloads/corporate-travel-policy-template.html is the source of
truth for wording. This writes the same document as a Word file so a financial
manager can edit it and forward it internally. Run after changing that HTML:

    python scripts/build_policy_docx.py
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = 'public/downloads/corporate-travel-policy-template.docx'
ACCENT = RGBColor(0x0F, 0x76, 0x6E)
MUTED = RGBColor(0x6B, 0x72, 0x80)
FILL = '_' * 30


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = ACCENT
    return p


def field_line(doc, label, width=30):
    p = doc.add_paragraph()
    p.add_run(label + ' ')
    p.add_run('_' * width)
    return p


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for cell, text in zip(t.rows[0].cells, headers):
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    return t


def note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)

title = doc.add_paragraph()
tr = title.add_run('Corporate Travel Policy Template')
tr.bold = True
tr.font.size = Pt(22)

sub = doc.add_paragraph()
sr = sub.add_run('A fillable starting point for South African businesses. Complete the fields '
                 'below, delete this note, and circulate internally. Prepared by RTM Travel '
                 '(Remmitz Travel Management) — rtmtravel.co.za')
sr.font.size = Pt(9)
sr.font.color.rgb = MUTED

table(doc, ['Company name', 'Effective date'], [['', '']])
table(doc, ['Policy owner', 'Next review date'], [['', '']])

heading(doc, '1. Advance Booking Requirements')
field_line(doc, 'Domestic flights and accommodation must be booked at least', 12)
doc.paragraphs[-1].add_run(' days before departure.')
field_line(doc, 'International flights and accommodation must be booked at least', 12)
doc.paragraphs[-1].add_run(' days before departure.')
note(doc, 'Standard: 14 days domestic, 21 days international.')

heading(doc, '2. Flight Class Guidelines')
table(doc, ['Trip type', 'Permitted class'], [
    ['Domestic (any duration)', ''],
    ['International, under ____ hours', ''],
    ['International, over ____ hours', ''],
    ['C-suite executives', ''],
])

heading(doc, '3. Accommodation Standards and Price Caps')
table(doc, ['City', 'Nightly rate cap (ZAR)'], [
    ['Johannesburg', ''],
    ['Cape Town', ''],
    ['Durban', ''],
    ['International (per city/region)', ''],
])

heading(doc, '4. Ground Transport Rules')
field_line(doc, 'Approved rental car category:')
field_line(doc, 'Approved ride-hailing services:')
field_line(doc, 'Executive car hire approval required from:')

heading(doc, '5. Meal and Per Diem Allowances')
table(doc, ['Travel type', 'Daily allowance (ZAR)'], [
    ['Domestic', ''],
    ['International', ''],
])

heading(doc, '6. Approval Workflows')
table(doc, ['Trip type', 'Approver'], [
    ['Standard domestic', ''],
    ['International', ''],
    ['Out-of-policy exception', ''],
])

heading(doc, '7. Duty of Care and Emergency Procedures')
field_line(doc, '24/7 emergency contact:')
field_line(doc, 'Emergency response protocol summary:', 60)
field_line(doc, 'Chain of communication (who is notified, in what order):', 45)

heading(doc, 'Sign-off')
table(doc, ['Role', 'Name', 'Signature', 'Date'], [
    ['Policy Owner', '', '', ''],
    ['Finance / CFO', '', '', ''],
])

foot = doc.add_paragraph()
foot.paragraph_format.space_before = Pt(24)
fr = foot.add_run('Template prepared by RTM Travel (rtmtravel.co.za) — a travel management '
                  'company can enforce every rule in this document automatically at the point '
                  'of booking. For the full explanation behind each section, see '
                  'https://rtmtravel.co.za/corporate-travel-policy-template')
fr.font.size = Pt(8.5)
fr.font.color.rgb = MUTED

doc.save(OUT)
print('wrote', OUT)
